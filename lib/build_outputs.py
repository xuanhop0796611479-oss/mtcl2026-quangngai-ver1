# -*- coding: utf-8 -*-
"""Dựng file Excel tổng hợp + dashboard HTML từ 2 file xlsx (ngày/tháng) + thống kê mất sóng."""
import os, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
HERE=os.path.dirname(os.path.abspath(__file__))
PROV="Quang Ngai"

def _rows(ws, hr, prov):
    H=[ws.cell(hr,c).value for c in range(1,47)]
    def ci(name):
        for i,h in enumerate(H):
            if h and name==str(h).strip(): return i
        for i,h in enumerate(H):
            if h and name in str(h): return i
        return None
    return H, ci

def extract_kpi(day_xlsx, month_xlsx):
    wb=openpyxl.load_workbook(day_xlsx, data_only=True); ws=wb.active; hr=3
    H,ci=_rows(ws,hr,PROV)
    keys={'xa':'Tỷ lệ Xã/phường đạt MTCL 2026','mtcl':'MTCL_2026','dg':'Đánh giá','kpi':'HTMT_KPI',
          'qos':'HTMT_QOS','vhkt':'HTMT_VHKT','mll':'MLL_TIME','traffic':'Traffic All (GB)','sotram':'Số lượng trạm'}
    ix={k:ci(v) for k,v in keys.items()}
    daily=[]
    for r in range(hr+1, ws.max_row+1):
        prov=ws.cell(r,4).value
        if prov and PROV in str(prov):
            v=[ws.cell(r,c).value for c in range(1,47)]; s=str(v[1])
            daily.append({'date':'%s/%s'%(s[6:8],s[4:6]),'xa':v[ix['xa']],'mtcl':v[ix['mtcl']],'dg':v[ix['dg']],
                'kpi':v[ix['kpi']],'qos':v[ix['qos']],'vhkt':v[ix['vhkt']],'mll':v[ix['mll']],'traffic':v[ix['traffic']]})
    daily.sort(key=lambda x:x['date'])
    wm=openpyxl.load_workbook(month_xlsx, data_only=True); wsm=wm.active; mr=4
    HM,cim=_rows(wsm,mr,PROV)
    mk={'xa':'Tỷ lệ Xã','mtcl':'MTCL_2026','dg':'Đánh giá','kpi':'HTMT_KPI','qos':'HTMT_QOS',
        'vhkt':'HTMT_VHKT','mll':'MLL_TIME','sotram':'Số lượng trạm','sl_xa':'SL Xã'}
    mix={k:cim(v) for k,v in mk.items()}
    monthly=None
    for r in range(mr+1, wsm.max_row+1):
        if wsm.cell(r,4).value and PROV in str(wsm.cell(r,4).value):
            v=[wsm.cell(r,c).value for c in range(1,47)]
            monthly={k:v[mix[k]] for k in mk}; break
    return daily, monthly

def build_excel(day_xlsx, month_xlsx, ms, out_path):
    thin=Side('thin', color='BFBFBF'); bd=Border(thin,thin,thin,thin)
    hf=PatternFill('solid', fgColor='1F4E79'); hff=Font(color='FFFFFF', bold=True, size=10)
    tf=Font(bold=True, size=12, color='1F4E79'); cf=Font(size=10)
    cen=Alignment('center','center',wrap_text=True); lef=Alignment('left','center')
    out=openpyxl.Workbook()
    def raw_sheet(ws_src, hr, title, name, wb):
        ws=wb.create_sheet(name) if name not in wb.sheetnames else wb[name]
        H=[ws_src.cell(hr,c).value for c in range(1,47)]
        cols=[i for i,h in enumerate(H) if h is not None]
        ws.append([title]); ws['A1'].font=tf; ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
        ws.append([H[i] for i in cols])
        for j in range(1,len(cols)+1):
            c=ws.cell(2,j); c.fill=hf; c.font=hff; c.alignment=cen; c.border=bd
        for r in range(hr+1, ws_src.max_row+1):
            prov=ws_src.cell(r,4).value
            if prov and PROV in str(prov):
                ws.append([ws_src.cell(r,i+1).value for i in cols]); rr=ws.max_row
                for j in range(1,len(cols)+1):
                    cc=ws.cell(rr,j); cc.font=cf; cc.border=bd; cc.alignment=lef if j<=4 else cen
        for j,i in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(j)].width=min(max(len(str(H[i]))+2,9),22)
        ws.freeze_panes=ws.cell(3,1)
    wbd=openpyxl.load_workbook(day_xlsx, data_only=True)
    ws1=out.active; ws1.title='KPI_Ngay_QuangNgai'
    raw_sheet(wbd.active, 3, 'MTCL2026 - Quảng Ngãi | KPI theo ngày', 'KPI_Ngay_QuangNgai', out)
    wbm=openpyxl.load_workbook(month_xlsx, data_only=True)
    raw_sheet(wbm.active, 4, 'MTCL2026 - Quảng Ngãi | Lũy kế tháng', 'LuyKe_Thang_QuangNgai', out)
    # MS sheets
    def ms_sheet(name, title, cols, rows):
        ws=out.create_sheet(name)
        ws.append([title]); ws['A1'].font=Font(bold=True,size=12,color='B71C1C'); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
        ws.append(cols)
        for j in range(1,len(cols)+1):
            c=ws.cell(2,j); c.fill=PatternFill('solid',fgColor='B71C1C'); c.font=hff; c.alignment=cen; c.border=bd
        for row in rows:
            ws.append(row); r=ws.max_row
            for j in range(1,len(cols)+1):
                c=ws.cell(r,j); c.font=cf; c.border=bd; c.alignment=lef if j<=3 else cen
        for j,cn in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(j)].width=min(max(len(str(cn))+2,10),34)
        ws.freeze_panes=ws.cell(3,1)
    ms_sheet('MatSong_TramKeoDai','Top trạm mất liên lạc kéo dài',
        ['Site','Lớp','Khu vực','Số phút','Trạng thái','Bắt đầu','Nguyên nhân'],
        [[x['site'],x['net'],x['district'],x['min'],'Đang mất' if x['ongoing'] else 'Đã KP',x['start'],x['cause'] or '-'] for x in ms['longest']])
    ms_sheet('MatSong_TanSuat','Top trạm mất sóng nhiều lần',
        ['Site','Khu vực','Lớp mạng','Số lần','Tổng phút','Dài nhất'],
        [[x['site'],x['district'],x['nets'],x['count'],x['total'],x['max']] for x in ms['freq']])
    ms_sheet('MatSong_DangMat','Trạm đang mất liên lạc',
        ['Site','Lớp','Khu vực','Số phút','Bắt đầu','Nguyên nhân'],
        [[x['site'],x['net'],x['district'],x['min'],x['start'],x['cause'] or '-'] for x in ms['ongoing_list']])
    out.save(out_path)

def build_dashboard(daily, monthly, ms, out_path, gen_ts=None, xa_fail=None, cell4g=None, xa_date=None, cell4g_date=None):
    import datetime
    if gen_ts is None:
        vn=datetime.datetime.now(datetime.timezone.utc)+datetime.timedelta(hours=7)
        gen_ts=vn.strftime('%d/%m/%Y %H:%M')

    payload={'daily':daily,'monthly':monthly,'ms':ms,
             'xa_fail':xa_fail or [], 'cell4g':cell4g or {},
             'xa_date':xa_date or '', 'cell4g_date':cell4g_date or ''}
    DATA=json.dumps(payload, ensure_ascii=False, default=str)
    tpl=open(os.path.join(HERE,'dashboard_template.html'), encoding='utf-8').read()
    js=open(os.path.join(HERE,'app_template.js'), encoding='utf-8').read()
    # Nội suy Chart.js để file tự chứa, chia sẻ được cả khi không có mạng
    chartjs_path=os.path.join(HERE,'vendor','chart.umd.min.js')
    chartjs=open(chartjs_path, encoding='utf-8').read() if os.path.exists(chartjs_path) else ''
    html=(tpl.replace('__CHARTJS__', chartjs)
             .replace('__GENTS__', gen_ts)
             .replace('__DATA__', DATA)
             .replace('<script src="app.js"></script>', '<script>\n'+js+'\n</script>'))
    open(out_path,'w',encoding='utf-8').write(html)


# ================== 2 BÁO CÁO BỞ SUNG ==================

def _to_num(x):
    if x is None or x=='':
        return None
    if isinstance(x,(int,float)):
        return float(x)
    s=str(x).strip().replace('\xa0','').replace(' ','')
    # xử lý số kiểu VN: dấu chấm phân cách nghìn, phẩy thập phân
    if s.count(',')==1 and s.count('.')>=1:
        s=s.replace('.','').replace(',','.')
    elif s.count(',')==1 and s.count('.')==0:
        s=s.replace(',','.')
    else:
        s=s.replace(',','')
    try:
        return float(s)
    except Exception:
        return None


import unicodedata as _ud
def _norm(s):
    """Bỏ dấu tiếng Việt + lowercase để so khớp tên cột không phụ thuộc dấu."""
    if s is None:
        return ''
    s=str(s).replace('\u0111','d').replace('\u0110','D')
    s=_ud.normalize('NFD', s)
    s=''.join(c for c in s if _ud.category(c)!='Mn')
    return s.strip().lower()


def _find_header(ws, must_have, max_scan=12):
    """Tìm dòng header chứa đủ các từ khóa must_have (không phụ thuộc dấu)."""


    ncol=min(ws.max_column, 80)
    for r in range(1, min(ws.max_row, max_scan)+1):
        vals=[ws.cell(r,c).value for c in range(1,ncol+1)]
        txt=[str(v).strip() if v is not None else '' for v in vals]
        low=[_norm(t) for t in txt]
        if all(any(_norm(k) in t for t in low) for k in must_have):
            return r, txt
    return None, None


def _col_index(header_txt, name):
    """Chỉ số cột (0-based) theo tên, không phụ thuộc dấu: khớp chính xác trước, rồi khớp chứa."""
    nm=_norm(name)
    for i,h in enumerate(header_txt):
        if h and nm==_norm(h):
            return i
    for i,h in enumerate(header_txt):
        if h and nm in _norm(h):
            return i
    return None


def _col_index_tokens(header_txt, tokens, avoid=None):
    """Chỉ số cột đầu tiên mà header (đã bỏ dấu) chứa ĐỦ mọi token trong `tokens`
    và KHÔNG chứa token nào trong `avoid`. Dùng để khớp tên cột kiểu
    'eUTraCell Name Old' / 'eUTrancell Old' / 'eNodeB Old' một cách linh hoạt."""
    toks=[_norm(t) for t in tokens]
    av=[_norm(t) for t in (avoid or [])]
    for i,h in enumerate(header_txt):
        if not h:
            continue
        hn=_norm(h)
        if all(t in hn for t in toks) and not any(a in hn for a in av):
            return i
    return None


def _first(*vals):
    """Trả về giá trị đầu tiên KHÁC None (chỉ số cột 0 vẫn hợp lệ)."""
    for v in vals:
        if v is not None:
            return v
    return None


def extract_xa_fail(xa_xlsx, province=PROV):
    """Trích danh sách XÃ/PHƯỜNG KHÔNG ĐẠT MTCL2026 (theo cột 'Đánh giá').
    Trả về (list_dicts, date_str)."""
    wb=openpyxl.load_workbook(xa_xlsx, data_only=True); ws=wb.active
    hr, H=_find_header(ws, ['MTCL', 'đánh giá'])
    if hr is None:
        hr, H=_find_header(ws, ['Xã', 'MTCL'])
    if hr is None:
        return [], ''
    ix={
        'xa':_first(_col_index(H,'Xã/Phường'), _col_index(H,'Xã')),
        'tinh':_col_index(H,'Tỉnh'),
        'ngay':_col_index(H,'Ngày'),
        'mtcl':_first(_col_index(H,'MTCL_2026'), _col_index(H,'MTCL')),
        'dg':_col_index(H,'Đánh giá'),
        'kpi':_col_index(H,'HTMT_KPI'),
        'qos':_col_index(H,'HTMT_QOS'),
        'vhkt':_col_index(H,'HTMT_VHKT'),
        'anm':_col_index(H,'HTMT_ANM'),
        'traffic':_first(_col_index(H,'TRAFFIC ALL'), _col_index(H,'Traffic All')),
    }
    def cell(row, key):
        i=ix.get(key)
        return row[i] if (i is not None and i<len(row)) else None
    out=[]; date_str=''
    for r in range(hr+1, ws.max_row+1):
        row=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if all(v is None or str(v).strip()=='' for v in row):
            continue
        # lọc tỉnh nếu có cột tỉnh
        if ix['tinh'] is not None:
            tv=cell(row,'tinh')
            if not tv or province not in str(tv):
                continue
        dg=cell(row,'dg'); dgn=_norm(dg).replace(' ','') if dg is not None else ''
        xa=cell(row,'xa')
        if not xa or str(xa).strip()=='':
            continue
        # chỉ giữ KHÔNG ĐẠT (dgn đã bỏ dấu: 'dat' = đạt)
        if dgn=='dat':
            continue
        if ix['ngay'] is not None and not date_str:
            nv=cell(row,'ngay')
            if nv: date_str=str(nv).strip()
        out.append({
            'xa':str(xa).strip(),
            'mtcl':_to_num(cell(row,'mtcl')),
            'dg':str(dg).strip() if dg is not None else '',
            'kpi':_to_num(cell(row,'kpi')),
            'qos':_to_num(cell(row,'qos')),
            'vhkt':_to_num(cell(row,'vhkt')),
            'anm':_to_num(cell(row,'anm')),
            'traffic':_to_num(cell(row,'traffic')),
        })
    out.sort(key=lambda x:(x['mtcl'] if x['mtcl'] is not None else 999))
    return out, date_str


def extract_4g_zero(g4_xlsx, province=PROV):
    """Trích danh sách cell 4G có TRAFFIC_4G (MB) = 0.
    Trả về (dict{'list':[...],'summary':{...}}, date_str)."""
    wb=openpyxl.load_workbook(g4_xlsx, data_only=True); ws=wb.active
    hr, H=_find_header(ws, ['TRAFFIC_4G'])
    if hr is None:
        hr, H=_find_header(ws, ['eUTrancell'])
    if hr is None:
        return {'list':[], 'summary':{'total':0,'zero':0}}, ''
    # Ưu tiên cột 'Old' theo yêu cầu (liệt kê eUTraCell Old). Cột 'New' thường
    # để trống nên nếu lấy New sẽ ra rỗng. Khớp linh hoạt cả 'eUTraCell' lẫn
    # 'eUTrancell', kể cả khi header có thêm chữ 'Name'.
    cell_ix=_first(_col_index_tokens(H, ['eutra','old']),
                   _col_index_tokens(H, ['eutra'], avoid=['new']),
                   _col_index_tokens(H, ['eutra','new']),
                   _col_index(H,'eUTrancell'))
    enb_ix=_first(_col_index_tokens(H, ['enodeb','old']),
                  _col_index_tokens(H, ['enodeb'], avoid=['new']),
                  _col_index_tokens(H, ['enodeb','new']),
                  _col_index(H,'eNodeB'))
    ix={
        'cell':cell_ix,
        'enb':enb_ix,
        'prov':_first(_col_index(H,'Province'), _col_index(H,'Tỉnh')),
        'xa':_first(_col_index(H,'Phường/Xã'), _col_index(H,'Xã/Phường'), _col_index(H,'Xã')),
        'ngay':_col_index(H,'Ngày'),
        'traffic':_first(_col_index(H,'TRAFFIC_4G'), _col_index_tokens(H, ['traffic','4g'])),
    }
    def cell(row, key):
        i=ix.get(key)
        return row[i] if (i is not None and i<len(row)) else None
    zero=[]; total=0; date_str=''
    for r in range(hr+1, ws.max_row+1):
        row=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if all(v is None or str(v).strip()=='' for v in row):
            continue
        cellname=cell(row,'cell')
        if not cellname or str(cellname).strip()=='':
            continue
        if ix['prov'] is not None:
            pv=cell(row,'prov')
            if pv and province not in str(pv):
                continue
        total+=1
        tr=_to_num(cell(row,'traffic'))
        if ix['ngay'] is not None and not date_str:
            nv=cell(row,'ngay')
            if nv: date_str=str(nv).strip()
        # TRAFFIC_4G (MB) = 0 (ô trống => coi như không có lưu lượng = 0)
        if tr is None or tr==0:
            zero.append({
                'cell':str(cellname).strip(),
                'enb':str(cell(row,'enb')).strip() if cell(row,'enb') is not None else '',
                'xa':str(cell(row,'xa')).strip() if cell(row,'xa') is not None else '',
                'traffic':0,
            })
    zero.sort(key=lambda x:x['cell'])
    summary={'total':total, 'zero':len(zero)}
    return {'list':zero, 'summary':summary}, date_str

